#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Compare an athlete registry XLSX with Athlete rows from a SQLite database."""

from __future__ import annotations

import argparse
import json
import re
import sqlite3
import subprocess
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


def normalize_name(value: Any) -> tuple[str, ...]:
    """Return an order-independent normalized FIO token key."""
    text = str(value or "").strip().lower().replace("ё", "е")
    text = re.sub(r"[^0-9a-zа-я-]+", " ", text)
    return tuple(sorted(part for part in text.split() if part))


def parse_registry_date(value: Any) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value).strip()
    for pattern in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, pattern).date()
        except ValueError:
            continue
    return None


def load_registry(path: Path) -> list[dict[str, Any]]:
    workbook = load_workbook(path, data_only=True, read_only=False)
    try:
        worksheet = workbook.active
        rows = []
        for row_number in range(2, worksheet.max_row + 1):
            fio_parts = [
                str(worksheet.cell(row_number, column).value or "").strip()
                for column in (1, 2, 3)
            ]
            fio = " ".join(part for part in fio_parts if part)
            if not fio:
                continue
            raw_birth_date = worksheet.cell(row_number, 4).value
            rows.append(
                {
                    "registry_row": row_number,
                    "fio": fio,
                    "name_key": normalize_name(fio),
                    "birth_date": parse_registry_date(raw_birth_date),
                    "birth_date_raw": raw_birth_date,
                    "organization": worksheet.cell(row_number, 5).value,
                    "rank": worksheet.cell(row_number, 7).value,
                }
            )
        return rows
    finally:
        workbook.close()


def load_database_rows(database_path: Path) -> list[dict[str, Any]]:
    connection = sqlite3.connect(database_path)
    connection.row_factory = sqlite3.Row
    try:
        existing_columns = {
            row["name"] for row in connection.execute("PRAGMA table_info(athlete)").fetchall()
        }
        optional_columns = [
            column
            for column in (
                "primary_external_id",
                "primary_first_name",
                "primary_last_name",
                "primary_patronymic",
                "primary_birth_date",
                "partner_external_id",
                "partner_first_name",
                "partner_last_name",
                "partner_patronymic",
                "partner_birth_date",
            )
            if column in existing_columns
        ]
        rows = connection.execute(
            "SELECT id, first_name, last_name, patronymic, full_name_xml, birth_date"
            + (", " + ", ".join(optional_columns) if optional_columns else "")
            + " FROM athlete"
        ).fetchall()
        return [dict(row) for row in rows]
    finally:
        connection.close()


def load_database_rows_over_ssh(
    host: str,
    port: int,
    user: str,
    key_path: Path,
    database_path: str,
) -> list[dict[str, Any]]:
    query = (
        "SELECT id, first_name, last_name, patronymic, full_name_xml, birth_date, "
        "primary_external_id, primary_first_name, primary_last_name, "
        "primary_patronymic, primary_birth_date, partner_external_id, "
        "partner_first_name, partner_last_name, partner_patronymic, "
        "partner_birth_date "
        "FROM athlete;"
    )
    command = [
        "ssh",
        "-i",
        str(key_path),
        "-p",
        str(port),
        "-o",
        "BatchMode=yes",
        f"{user}@{host}",
        f"sqlite3 -json {database_path!s} {json.dumps(query)}",
    ]
    completed = subprocess.run(
        command,
        check=True,
        capture_output=True,
        text=True,
        encoding="utf-8",
    )
    return json.loads(completed.stdout or "[]")


def database_name(row: dict[str, Any]) -> str:
    full_name = str(row.get("full_name_xml") or "").strip()
    if full_name:
        return full_name
    return " ".join(
        str(row.get(field) or "").strip()
        for field in ("last_name", "first_name", "patronymic")
        if str(row.get(field) or "").strip()
    )


def compare(
    registry_rows: list[dict[str, Any]],
    database_rows: list[dict[str, Any]],
) -> dict[str, list[dict[str, Any]]]:
    database_by_name: dict[tuple[str, ...], list[dict[str, Any]]] = defaultdict(list)
    for row in database_rows:
        fio = database_name(row)
        if "/" not in fio:
            individual = {**row, "fio": fio, "name_key": normalize_name(fio)}
            database_by_name[individual["name_key"]].append(individual)
            continue

        # A pair is one result record, but both people can be matched against
        # the external athlete registry after the pair-member migration.
        for prefix in ("primary", "partner"):
            member_fio = " ".join(
                str(row.get(f"{prefix}_{field}") or "").strip()
                for field in ("last_name", "first_name", "patronymic")
                if str(row.get(f"{prefix}_{field}") or "").strip()
            )
            if not member_fio:
                continue
            member = {
                **row,
                "fio": member_fio,
                "birth_date": row.get(f"{prefix}_birth_date"),
                "name_key": normalize_name(member_fio),
            }
            database_by_name[member["name_key"]].append(member)

    result: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for registry in registry_rows:
        candidates = database_by_name.get(registry["name_key"], [])
        base = {
            "registry_row": registry["registry_row"],
            "registry_fio": registry["fio"],
            "registry_birth_date": registry["birth_date"],
            "organization": registry["organization"],
            "rank": registry["rank"],
        }
        if not candidates:
            result["not_found"].append(base)
            continue

        if len(candidates) > 1:
            same_date = [
                candidate
                for candidate in candidates
                if registry["birth_date"]
                and parse_registry_date(candidate.get("birth_date")) == registry["birth_date"]
            ]
            if len(same_date) == 1:
                candidates = same_date
            else:
                result["ambiguous"].append(
                    {
                        **base,
                        "database_candidates": "; ".join(
                            f"ID {candidate['id']}: {candidate['fio']} "
                            f"({candidate.get('birth_date') or 'без ДР'})"
                            for candidate in candidates
                        ),
                    }
                )
                continue

        candidate = candidates[0]
        database_birth_date = parse_registry_date(candidate.get("birth_date"))
        combined = {
            **base,
            "database_id": candidate["id"],
            "database_fio": candidate["fio"],
            "database_birth_date": database_birth_date,
        }
        if registry["birth_date"] and database_birth_date:
            bucket = "same_date" if registry["birth_date"] == database_birth_date else "different_date"
        elif registry["birth_date"] and not database_birth_date:
            bucket = "missing_in_database"
        elif not registry["birth_date"] and database_birth_date:
            bucket = "missing_in_registry"
        else:
            bucket = "both_dates_missing"
        result[bucket].append(combined)
    return result


def display_date(value: Any) -> str:
    parsed = parse_registry_date(value)
    return parsed.strftime("%d.%m.%Y") if parsed else ""


def write_sheet(
    workbook: Workbook,
    title: str,
    rows: list[dict[str, Any]],
    columns: list[tuple[str, str]],
) -> None:
    worksheet = workbook.create_sheet(title)
    worksheet.append([header for _, header in columns])
    for cell in worksheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
    for row in rows:
        worksheet.append(
            [
                display_date(row.get(key)) if "birth_date" in key else row.get(key, "")
                for key, _ in columns
            ]
        )
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for index, (_, header) in enumerate(columns, start=1):
        values = [str(worksheet.cell(row, index).value or "") for row in range(1, worksheet.max_row + 1)]
        worksheet.column_dimensions[get_column_letter(index)].width = min(
            max(len(header), *(len(value) for value in values)) + 2,
            60,
        )


def write_report(path: Path, result: dict[str, list[dict[str, Any]]]) -> None:
    workbook = Workbook()
    workbook.remove(workbook.active)

    summary = workbook.create_sheet("Сводка")
    summary.append(["Категория", "Количество"])
    labels = [
        ("different_date", "Расхождение даты"),
        ("missing_in_database", "Дата отсутствует в БД"),
        ("missing_in_registry", "Дата отсутствует в реестре"),
        ("ambiguous", "Неоднозначное ФИО"),
        ("not_found", "Не найдено в БД"),
        ("same_date", "ФИО и дата совпадают"),
        ("both_dates_missing", "Обе даты отсутствуют"),
    ]
    for key, label in labels:
        summary.append([label, len(result.get(key, []))])
    summary.column_dimensions["A"].width = 34
    summary.column_dimensions["B"].width = 14
    for cell in summary[1]:
        cell.font = Font(bold=True)

    matched_columns = [
        ("registry_row", "Строка Excel"),
        ("registry_fio", "ФИО в реестре"),
        ("registry_birth_date", "Дата в реестре"),
        ("database_id", "ID в БД"),
        ("database_fio", "ФИО в БД"),
        ("database_birth_date", "Дата в БД"),
        ("organization", "Организация"),
        ("rank", "Разряд"),
    ]
    write_sheet(workbook, "Расхождения дат", result.get("different_date", []), matched_columns)
    write_sheet(workbook, "Нет даты в БД", result.get("missing_in_database", []), matched_columns)
    write_sheet(workbook, "Нет даты в Excel", result.get("missing_in_registry", []), matched_columns)
    write_sheet(workbook, "Совпадают", result.get("same_date", []), matched_columns)
    write_sheet(workbook, "Обе даты пустые", result.get("both_dates_missing", []), matched_columns)
    write_sheet(
        workbook,
        "Неоднозначные",
        result.get("ambiguous", []),
        [
            ("registry_row", "Строка Excel"),
            ("registry_fio", "ФИО в реестре"),
            ("registry_birth_date", "Дата в реестре"),
            ("database_candidates", "Кандидаты в БД"),
            ("organization", "Организация"),
        ],
    )
    write_sheet(
        workbook,
        "Не найдены",
        result.get("not_found", []),
        [
            ("registry_row", "Строка Excel"),
            ("registry_fio", "ФИО в реестре"),
            ("registry_birth_date", "Дата в реестре"),
            ("organization", "Организация"),
            ("rank", "Разряд"),
        ],
    )
    workbook.save(path)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--registry", required=True, type=Path)
    parser.add_argument("--output", required=True, type=Path)
    source = parser.add_mutually_exclusive_group(required=True)
    source.add_argument("--database", type=Path)
    source.add_argument("--ssh-host")
    parser.add_argument("--ssh-port", type=int, default=2222)
    parser.add_argument("--ssh-user", default="root")
    parser.add_argument("--ssh-key", type=Path)
    parser.add_argument(
        "--remote-database",
        default="/var/www/calc.figurebase.ru/instance/figure_skating.db",
    )
    args = parser.parse_args()

    registry_rows = load_registry(args.registry)
    if args.database:
        database_rows = load_database_rows(args.database)
    else:
        if not args.ssh_key:
            parser.error("--ssh-key is required with --ssh-host")
        database_rows = load_database_rows_over_ssh(
            args.ssh_host,
            args.ssh_port,
            args.ssh_user,
            args.ssh_key,
            args.remote_database,
        )

    result = compare(registry_rows, database_rows)
    write_report(args.output, result)
    print(f"Registry rows: {len(registry_rows)}")
    print(f"Database rows: {len(database_rows)}")
    for key in (
        "different_date",
        "missing_in_database",
        "missing_in_registry",
        "ambiguous",
        "not_found",
        "same_date",
        "both_dates_missing",
    ):
        print(f"{key}: {len(result.get(key, []))}")
    print(f"Report: {args.output}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
